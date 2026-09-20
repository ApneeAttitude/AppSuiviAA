#!/usr/bin/env python3
"""Contrôle de non-régression sur les classeurs générés."""
import glob
import sys
import openpyxl

from build_suivi import CREF_FIRST, ROSTER_N

ERR = 0


def fail(msg):
    global ERR
    ERR += 1
    print("  ÉCHEC :", msg)


def check_erreurs(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("#"):
                    bad.append("%s!%s = %s" % (ws.title, c.coordinate, c.value))
    if bad:
        fail("%d cellule(s) en erreur : %s" % (len(bad), bad[:5]))
    return wb


def main():
    par = "../classeurs/AA - Parametrage 2026-2027.xlsx"
    print("paramétrage :", par)
    wb = check_erreurs(par)

    ws = wb["Objectifs_Niveaux"]
    n = sum(1 for r in range(5, 60) if ws.cell(r, 1).value)
    if n != 46:
        fail("46 objectifs attendus, %d trouvés" % n)
    seuils = {}
    for r in range(5, 60):
        g, t, s = ws.cell(r, 2).value, ws.cell(r, 4).value, ws.cell(r, 12).value
        if g and t == "maîtrise":
            seuils[g] = s
    for g, attendu in (("L1", 7), ("L2", 5), ("L3", 3)):
        if seuils.get(g) != attendu:
            fail("seuil de maîtrise %s : %s attendu, %s trouvé"
                 % (g, attendu, seuils.get(g)))
    if seuils.get("L4") not in (None, ""):
        fail("le seuil L4 doit rester vide (critère qualitatif)")

    ws = wb["Personnes"]
    npers = sum(1 for r in range(5, 205) if ws.cell(r, 1).value)
    if npers < 60:
        fail("moins de 60 personnes dans le référentiel (%d)" % npers)
    if not ws.cell(5, 10).value:
        fail("libellés de sélection du paramétrage non calculés")
    sel = wb["Selection"]
    n_sel = sum(1 for r in range(2, 202) if sel.cell(r, 1).value)
    if n_sel != npers:
        fail("liste de sélection : %d entrées pour %d personnes"
             % (n_sel, npers))
    for nom in ("Referentiels", "Competences_Niveaux",
                "Objectifs_Niveaux", "Creneaux", "Creneaux_Affectations",
                "Lignes_Seances", "Regles_Acces_Seances"):
        if nom not in wb.sheetnames:
            fail("onglet %s absent" % nom)
    if not wb["Referentiels"].cell(53, 1).value:
        fail("section Familles de compétences vide dans Referentiels (ligne 53)")
    cg = wb["Creneaux_Affectations"]
    if not cg.cell(5, 2).value:
        fail("identifiant de créneau non déduit dans Creneaux_Affectations")

    ws = wb["Inscriptions"]
    ninsc = sum(1 for r in range(5, 405)
                if ws.cell(r, 3).value not in (None, ""))
    if ninsc == 0:
        print("  info : aucune inscription saisie — les classeurs de ligne "
              "resteront vides jusqu'à la saisie")
    ctrl = [ws.cell(r, 13).value for r in range(5, 405)
            if ws.cell(r, 3).value not in (None, "")]
    ko = [c for c in ctrl if c != "ok"]
    if ko:
        print("  info : %d inscription(s) en anomalie (%s)"
              % (len(ko), sorted(set(map(str, ko)))))

    if "Responsables_Ligne" not in wb.sheetnames:
        fail("onglet Responsables_Ligne absent")
    else:
        ws = wb["Responsables_Ligne"]
        nres = sum(1 for r in range(5, 105)
                   if ws.cell(r, 4).value not in (None, ""))
        if nres == 0:
            print("  info : aucun responsable de ligne saisi — la couverture "
                  "des séances restera incomplète")
    print("  46 objectifs · seuils 7/5/3/qualitatif · %d personnes · "
          "%d inscriptions" % (npers, ninsc))

    attendus = {"Lisez-moi", "Synchro", "Referentiel", "Personnes",
                "Calendrier", "Presences",
                "Suivi_Objectifs", "Suivi_Competences", "Avis_Encadrants",
                "Evenements_Securite", "Zones_Physiologiques",
                "Tableau_de_bord", "Selection", "Listes"}

    for path in sorted(glob.glob("../classeurs/AA - Suivi *.xlsx")):
        print("ligne :", path.split("/")[-1])
        wb = check_erreurs(path)
        manquants = attendus - set(wb.sheetnames)
        if manquants:
            fail("onglets manquants : %s" % sorted(manquants))
            continue
        if not wb["Synchro"].cell(9, 2).value:
            fail("classeur non synchronisé (onglet Synchro incomplet)")
        ref = wb["Referentiel"]
        if not ref.cell(6, 1).value:
            fail("référentiel d'objectifs vide (attendu en ligne 6)")
        if not ref.cell(30, 1).value:
            fail("référentiel de compétences vide (attendu en ligne 30)")
        if not str(ref.cell(4, 1).value or "").startswith("1."):
            fail("titre du bloc objectifs absent")
        if not str(ref.cell(28, 1).value or "").startswith("2."):
            fail("titre du bloc compétences absent")
        if ref.cell(CREF_FIRST, 1).value and not ref.cell(CREF_FIRST, 7).value:
            fail("libellé de créneau non calculé")
        pers = wb["Personnes"]
        membres = sum(1 for r in range(5, 155) if pers.cell(r, 8).value == "oui")
        eleves = sum(1 for r in range(5, 155) if pers.cell(r, 12).value not in
                     (None, ""))
        codes = [pers.cell(r, 1).value for r in range(5, 155)
                 if pers.cell(r, 1).value]
        if membres and codes != list(range(1, membres + 1)):
            fail("codes des membres non séquentiels : %s" % codes[:10])
        # le roster des matrices (élèves seulement) doit se remplir tout seul
        for onglet, first in (("Suivi_Objectifs", 8), ("Suivi_Competences", 6),
                              ("Avis_Encadrants", 10)):
            wsx = wb[onglet]
            remplis = sum(1 for i in range(ROSTER_N)
                          if wsx.cell(first + i, 2).value not in (None, ""))
            if remplis != min(eleves, ROSTER_N):
                fail("%s : %d apnéistes construits, %d élèves attendus"
                     % (onglet, remplis, min(eleves, ROSTER_N)))
        pre = wb["Presences"]
        if "Présent" in [pre.cell(4, c).value for c in range(1, 12)]:
            fail("la colonne « Présent » ne devrait plus exister")
        cal = wb["Calendrier"]
        entetes = [cal.cell(4, c).value for c in range(1, 22)]
        if not any(str(x) == "Responsable remplaçant — début du nom" for x in entetes):
            fail("colonne de remplacement du responsable absente")
        print("  %d membres · roster automatique OK · %d objectifs au "
              "référentiel" % (membres, wb["Synchro"].cell(10, 2).value or 0))

    print("\n%s" % ("CONTRÔLE OK" if ERR == 0 else "%d ÉCHEC(S)" % ERR))
    sys.exit(1 if ERR else 0)


if __name__ == "__main__":
    main()
